import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_files(estimate_data, base_name="estimate"):
    """
    Принимает массив словарей со сметой и генерирует два файла: .txt и .xlsx.
    Возвращает кортеж путей к файлам: (txt_path, xlsx_path).
    """
    # Создаем временную директорию для сборки файлов, если её нет
    os.makedirs("temp", exist_ok=True)
    
    txt_path = f"temp/{base_name}.txt"
    xlsx_path = f"temp/{base_name}.xlsx"
    
    # --- 1. ГЕНЕРАЦИЯ TXT ФАЙЛА ---
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("СТРОИТЕЛЬНАЯ СМЕТА НА ОСНОВЕ ДИЗАЙН-ПРОЕКТА\n")
        f.write("="*60 + "\n\n")
        
        grand_total = 0
        for item in estimate_data:
            name = item.get("name", "Неизвестная работа")
            unit = item.get("unit", "ед.")
            count = item.get("count", 0)
            price = item.get("price", 0)
            total = item.get("total", 0)
            grand_total += total
            
            f.write(f"Наименование: {name}\n")
            f.write(f"Объем: {count} {unit} | Цена: {price} руб. | Итого: {total} руб.\n")
            f.write("-" * 40 + "\n")
            
        f.write(f"\nИТОГОВАЯ СТОИМОСТЬ РАБОТ: {grand_total} руб.\n")

    # --- 2. ГЕНЕРАЦИЯ EXCEL ФАЙЛА (.xlsx) ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Смета"
    
    # Включаем сетку таблицы (чтобы на телефоне в Excel было всё видно)
    ws.views.sheetView[0].showGridLines = True
    
    # Шапка таблицы
    headers = ["Наименование работ", "Ед. изм.", "Количество", "Цена (руб.)", "Стоимость (руб.)"]
    ws.append(headers)
    
    # Стилизация шапки (темно-серый профессиональный стиль)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for col_num in range(1, 6):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        
    # Стили для строк с данными
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    # Заполняем строками
    for item in estimate_data:
        row_data = [
            item.get("name", ""),
            item.get("unit", ""),
            item.get("count", 0),
            item.get("price", 0),
            item.get("total", 0)
        ]
        ws.append(row_data)
        
        # Применяем границы к последней добавленной строке
        curr_row = ws.max_row
        for col_num in range(1, 6):
            c = ws.cell(row=curr_row, column=col_num)
            c.border = thin_border
            if col_num in [2, 3, 4, 5]:
                c.alignment = Alignment(horizontal="center")

    # Добавляем финальную строчку "Итого"
    last_row = ws.max_row + 2
    ws.cell(row=last_row, column=1, value="ИТОГО ПО СМЕТЕ:").font = Font(bold=True)
    # Используем формулу Excel для подсчета суммы столбца E
    ws.cell(row=last_row, column=5, value=f"=SUM(E2:E{last_row-2})").font = Font(bold=True)
    
    # Автоподбор ширины колонок
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    wb.save(xlsx_path)
    
    return txt_path, xlsx_path
  
